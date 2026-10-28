"""
generate_weekly_report.py — build and email a detailed weekly cleanup
report summarizing what cleaning.py's reconciliation changed, plus items
that need human review.

Reads the changelog.json that cleaning.py writes to S3, builds a chart
(what changed, embedded in the Word doc), a full-detail Excel workbook,
and a narrative Word document, then emails the Excel + Word doc via SES
with a highlights summary in the email body itself. Run after cleaning.py
for the same --date.

Usage:
    python generate_weekly_report.py --date 2026-07-28
    python generate_weekly_report.py --date 2026-07-28 --dry-run   # build files, don't send
"""

import argparse
import os
import shutil
import tempfile
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import boto3
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import polars as pl
import xlsxwriter
from docx import Document
from docx.shared import Inches

from db import Database
from report_brand_duplicates import _containment_pairs, _load_brand_counts, _similarity_pairs
from s3 import S3Client

S3_BUCKET = "raw-data-prod"


def load_changelog(s3: S3Client, run_date: str) -> dict:
    key = f"beer-pricing/cleanup-reports/{run_date}/changelog.json"
    return s3.read_json(f"s3://{S3_BUCKET}/{key}")


def load_duplicate_pairs(db: Database) -> list[dict]:
    counts = _load_brand_counts(db)
    brands = sorted(counts)
    containment = _containment_pairs(brands)
    similarity = [
        pair for pair in _similarity_pairs(brands)
        if pair not in containment and pair[::-1] not in containment
    ]
    return [
        {"brand_a": a, "count_a": counts[a], "brand_b": b, "count_b": counts[b], "kind": kind}
        for kind, pairs in (("brand/sub-brand", containment), ("near-identical spelling", similarity))
        for a, b in pairs
    ]


def build_chart(changelog: dict, out_path: str) -> None:
    labels = ["New", "Confirmed", "Reset", "Other segment", "Suspicious links", "NA/invalid skipped"]
    values = [
        len(changelog["new_links"]),
        len(changelog["confirmed_links"]),
        len(changelog["reset_links"]),
        sum(row["count"] for row in changelog["other_segment_brands"]),
        len(changelog["suspicious_links"]),
        len(changelog.get("invalid_links", [])),
    ]
    fig, ax = plt.subplots(figsize=(8, 4.5))
    bars = ax.bar(labels, values, color="#3b6ea5")
    ax.set_title(f"Weekly cleanup summary — {changelog['date']}")
    ax.set_ylabel("Count")
    for bar, value in zip(bars, values):
        ax.annotate(str(value), (bar.get_x() + bar.get_width() / 2, bar.get_height()), ha="center", va="bottom")
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)


def _rows_to_df(rows: list[dict], fallback_cols: list[str]) -> pl.DataFrame:
    if not rows:
        return pl.DataFrame({col: [] for col in fallback_cols})
    return pl.DataFrame(rows)


# Fields worth a human glance when placeholder — quantity/pack_format/units
# being NA is a real extraction gap; platform/link/currency being NA never
# happens (those are set by the pipeline itself, not extracted per-product).
_QA_COLUMNS = [
    "product_name", "brand", "cleaned_brand", "type", "sub_type", "quantity",
    "units", "pack_format", "original_price_local", "discounted_price_local",
    "availability", "region", "segment", "country",
]


_PRICE_COLUMNS = ("original_price_local", "discounted_price_local")


def _is_placeholder(value) -> bool:
    if value is None:
        return True
    text = str(value).strip()
    return text == "" or text.upper() in ("NA", "N/A", "-", "--")


def load_unverified_rows(db: Database, run_date: str) -> pl.DataFrame:
    """Every row this week with at least one placeholder field (NULL/""/NA/-)
    in any QA-relevant column — a one-page checklist for reviewing what the
    extraction pipeline couldn't fill in, instead of hunting for dashes
    scattered across the full per-platform data."""
    cols = ["link", "platform"] + _QA_COLUMNS
    df = db.read_sql(
        f"SELECT {', '.join(cols)} FROM beer_pricing_analysis.\"plaformsV2\" WHERE date = '{run_date}'",
        override_db="raw_data",
    )

    flagged = []
    for row in df.to_dicts():
        # Out-of-stock items legitimately have no price — not an extraction gap.
        out_of_stock = str(row.get("availability")).strip().lower() == "no"
        bad_fields = [
            col for col in _QA_COLUMNS
            if _is_placeholder(row.get(col))
            and not (out_of_stock and col in _PRICE_COLUMNS)
        ]
        if bad_fields:
            flagged.append({
                "link": row["link"],
                "platform": row["platform"],
                "flagged_fields": ", ".join(bad_fields),
                **{col: row.get(col) for col in _QA_COLUMNS},
            })

    return _rows_to_df(flagged, ["link", "platform", "flagged_fields"] + _QA_COLUMNS)


def build_excel(changelog: dict, duplicate_pairs: list[dict], unverified: pl.DataFrame, out_path: str) -> None:
    summary = pl.DataFrame(
        [
            {"metric": "Total rows written", "value": changelog["total_rows"]},
            {"metric": "New pending links", "value": len(changelog["new_links"])},
            {"metric": "Newly confirmed links", "value": len(changelog["confirmed_links"])},
            {"metric": "Reset (disagreement) links", "value": len(changelog["reset_links"])},
            {"metric": "Pending links that agreed again", "value": changelog["incremented_count"]},
            {"metric": "Suspicious shared-link flags", "value": len(changelog["suspicious_links"])},
            {"metric": "Links skipped (NA/invalid field this run)", "value": len(changelog.get("invalid_links", []))},
            {
                "metric": "Rows in 'other' segment",
                "value": sum(row["count"] for row in changelog["other_segment_brands"]),
            },
            {"metric": "Rows with an unverified/placeholder field", "value": unverified.height},
        ]
    )
    rows_per_platform = pl.DataFrame(
        [{"platform": k, "rows": v} for k, v in changelog["rows_per_platform"].items()]
    ).sort("platform")

    sheets = [
        ("Summary", summary),
        ("Rows Per Platform", rows_per_platform),
        ("Reconciliation Changes", _rows_to_df(changelog["reconciliation_changes"], ["link", "column", "before", "after"])),
        ("New Products", _rows_to_df(changelog["new_links"], ["link", "cleaned_brand", "sub_type", "platform", "brand", "pack_format"])),
        ("Newly Confirmed", _rows_to_df(changelog["confirmed_links"], ["link"])),
        ("Disagreements-Reset", _rows_to_df(changelog["reset_links"], ["link"])),
        ("Near-Duplicate Brands", _rows_to_df(duplicate_pairs, ["brand_a", "count_a", "brand_b", "count_b", "kind"])),
        ("Shared-Link Flags", pl.DataFrame({"link": changelog["suspicious_links"]})),
        ("NA-Invalid Skipped", _rows_to_df(changelog.get("invalid_links", []), ["link", "cleaned_brand", "sub_type", "platform", "brand", "pack_format"])),
        ("Other Segment Brands", _rows_to_df(changelog["other_segment_brands"], ["cleaned_brand", "count"])),
        ("Unverified-Errors", unverified),
    ]

    workbook = xlsxwriter.Workbook(out_path)
    for name, frame in sheets:
        frame.write_excel(workbook=workbook, worksheet=name, autofit=True)
    workbook.close()


def build_docx(changelog: dict, duplicate_pairs: list[dict], chart_path: str, out_path: str) -> None:
    doc = Document()
    doc.add_heading(f"Weekly Cleanup Report — {changelog['date']}", level=1)

    doc.add_paragraph(
        f"This run wrote {changelog['total_rows']} row(s) to beer_pricing_analysis.plaformsV2. "
        f"{len(changelog['new_links'])} new product link(s) were seen for the first time, "
        f"{len(changelog['confirmed_links'])} link(s) became trusted (confirmed) this week, and "
        f"{len(changelog['reset_links'])} pending link(s) disagreed with their prior candidate "
        "value and had their confirmation streak reset."
    )

    doc.add_picture(chart_path, width=Inches(6))

    doc.add_heading("Needs Human Review", level=2)

    doc.add_heading("Shared/placeholder link flags (high priority)", level=3)
    if changelog["suspicious_links"]:
        doc.add_paragraph(
            "These links were seen with several distinct brands in a single run — almost "
            "certainly a shared/placeholder URL (e.g. a WhatsApp contact link) rather than a "
            "real per-product page. products_source was NOT updated for these."
        )
        for link in changelog["suspicious_links"]:
            doc.add_paragraph(link, style="List Bullet")
    else:
        doc.add_paragraph("None this run.")

    doc.add_heading("Links that disagreed with their prior candidate", level=3)
    if changelog["reset_links"]:
        for row in changelog["reset_links"]:
            doc.add_paragraph(f"{row['link']} -> brand={row.get('brand')}", style="List Bullet")
    else:
        doc.add_paragraph("None this run.")

    doc.add_heading("Links skipped for NA/invalid field data", level=3)
    invalid_links = changelog.get("invalid_links", [])
    if invalid_links:
        doc.add_paragraph(
            "These links had a missing/NA value in a reconciled field this run (e.g. an "
            "extraction failure on the source page). They are never frozen into "
            "products_source while this keeps happening — agreement between weeks is not "
            "the same as a correct value, so NA never counts as \"confirmed.\" If a link "
            "stays here for multiple weeks running, the underlying scraper/extractor likely "
            "needs a fix."
        )
        for row in invalid_links:
            na_fields = [c for c in ("cleaned_brand", "sub_type", "platform", "brand", "pack_format") if not row.get(c) or str(row.get(c)).strip().upper() in ("NA", "N/A")]
            doc.add_paragraph(f"{row['link']} -> NA in: {', '.join(na_fields)}", style="List Bullet")
    else:
        doc.add_paragraph("None this run.")

    doc.add_heading("Possible duplicate brand spellings", level=3)
    if duplicate_pairs:
        for pair in duplicate_pairs:
            doc.add_paragraph(
                f"{pair['brand_a']!r} (n={pair['count_a']})  <->  {pair['brand_b']!r} "
                f"(n={pair['count_b']})  [{pair['kind']}]",
                style="List Bullet",
            )
    else:
        doc.add_paragraph("None this run.")

    doc.add_heading("Newly Confirmed Products", level=2)
    if changelog["confirmed_links"]:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        header = table.rows[0].cells
        header[0].text, header[1].text, header[2].text = "Link", "Brand", "Pack Format"
        for row in changelog["confirmed_links"]:
            cells = table.add_row().cells
            cells[0].text = str(row.get("link"))
            cells[1].text = str(row.get("brand"))
            cells[2].text = str(row.get("pack_format"))
    else:
        doc.add_paragraph("None this run.")

    doc.add_paragraph(
        "\nFull row-level detail for every changed record this run is in the attached Excel workbook."
    )

    doc.save(out_path)


def _build_highlights_html(changelog: dict, duplicate_pairs: list[dict], unverified: pl.DataFrame, run_date: str) -> str:
    new_count = len(changelog["new_links"])
    confirmed_count = len(changelog["confirmed_links"])
    reset_count = len(changelog["reset_links"])
    suspicious_count = len(changelog["suspicious_links"])
    invalid_count = len(changelog.get("invalid_links", []))
    other_count = sum(row["count"] for row in changelog["other_segment_brands"])
    dup_count = len(duplicate_pairs)
    unverified_count = unverified.height
    platform_count = len(changelog["rows_per_platform"])

    top_platforms = sorted(changelog["rows_per_platform"].items(), key=lambda kv: kv[1], reverse=True)[:5]
    top_platforms_html = "".join(f"<li>{name}: {count}</li>" for name, count in top_platforms)

    attention_items = []
    if suspicious_count:
        attention_items.append(
            f"<li><strong>{suspicious_count} shared/placeholder link flag(s)</strong> — a single link "
            "was seen with several distinct brands this run, most likely a non-unique URL. "
            "products_source was <em>not</em> updated for these. See \"Shared-Link Flags\" in the Excel.</li>"
        )
    if reset_count:
        attention_items.append(
            f"<li><strong>{reset_count} pending link(s) disagreed</strong> with their prior candidate value "
            "and had their confirmation streak reset. See \"Disagreements-Reset\" in the Excel.</li>"
        )
    if dup_count:
        attention_items.append(
            f"<li><strong>{dup_count} possible duplicate brand spelling(s)</strong> found — review "
            "\"Near-Duplicate Brands\" in the Excel.</li>"
        )
    if invalid_count:
        attention_items.append(
            f"<li><strong>{invalid_count} link(s) skipped</strong> for having an NA/missing field value "
            "this run — never frozen into products_source while this keeps happening. Links "
            "appearing here repeatedly likely need a scraper/extractor fix. See \"NA-Invalid "
            "Skipped\" in the Excel.</li>"
        )
    if unverified_count:
        attention_items.append(
            f"<li><strong>{unverified_count} row(s) have an unverified/placeholder field</strong> "
            "(blank, \"-\", or NA in brand, sub_type, quantity, units, pack_format, or similar) — "
            "a quick QA checklist, not necessarily errors. See \"Unverified-Errors\" in the Excel.</li>"
        )
    attention_html = (
        f"<ul>{''.join(attention_items)}</ul>" if attention_items else "<p>Nothing needs attention this week.</p>"
    )

    return f"""
    <html><body style="font-family: -apple-system, Segoe UI, Arial, sans-serif; color: #1f2933;">
    <h2>Weekly Cleanup Report — {run_date}</h2>
    <p>{changelog['total_rows']} row(s) written to <code>plaformsV2</code> across {platform_count} platform(s).</p>

    <h3>Highlights</h3>
    <ul>
      <li><strong>{new_count}</strong> new product link(s) seen for the first time</li>
      <li><strong>{confirmed_count}</strong> link(s) newly confirmed (trusted) this week</li>
      <li><strong>{reset_count}</strong> link(s) disagreed with their prior candidate and were reset</li>
      <li><strong>{other_count}</strong> row(s) fell into the uncategorized "other" segment</li>
      <li><strong>{suspicious_count}</strong> shared/placeholder link flag(s)</li>
      <li><strong>{invalid_count}</strong> link(s) skipped for NA/invalid field data</li>
      <li><strong>{dup_count}</strong> possible duplicate brand spelling(s) to review</li>
      <li><strong>{unverified_count}</strong> row(s) have an unverified/placeholder field this week</li>
    </ul>

    <h3>Top platforms by row count</h3>
    <ul>{top_platforms_html}</ul>

    <h3>Needs attention</h3>
    {attention_html}

    <p>Full record-level detail for every changed record is in the attached Excel workbook; a
    narrative summary (including this run's summary chart) is in the attached Word document.</p>
    </body></html>
    """


def send_report_email(
    xlsx_path: str,
    docx_path: str,
    changelog: dict,
    duplicate_pairs: list[dict],
    unverified: pl.DataFrame,
    run_date: str,
    sender: str,
    recipients: list[str],
) -> None:
    msg = MIMEMultipart("mixed")
    msg["Subject"] = f"Price Scraper — Weekly Cleanup Report — {run_date}"
    msg["From"] = sender
    msg["To"] = ", ".join(recipients)

    html = _build_highlights_html(changelog, duplicate_pairs, unverified, run_date)
    msg.attach(MIMEText(html, "html"))

    for path, filename in [
        (xlsx_path, f"cleanup_report_{run_date}.xlsx"),
        (docx_path, f"cleanup_report_{run_date}.docx"),
    ]:
        with open(path, "rb") as f:
            part = MIMEApplication(f.read())
        part.add_header("Content-Disposition", "attachment", filename=filename)
        msg.attach(part)

    ses = boto3.client("ses", region_name="us-east-1")
    ses.send_raw_email(Source=sender, Destinations=recipients, RawMessage={"Data": msg.as_string()})
    print(f"Sent weekly report email to {', '.join(recipients)}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build and email the weekly cleanup report.")
    parser.add_argument("--date", required=True, help="Run date (YYYY-MM-DD), matching cleaning.py's --date.")
    parser.add_argument(
        "--dry-run", action="store_true", help="Build the report files locally without sending an email."
    )
    args = parser.parse_args()

    s3 = S3Client()
    db = Database("raw_data")

    changelog = load_changelog(s3, args.date)
    duplicate_pairs = load_duplicate_pairs(db)
    unverified = load_unverified_rows(db, args.date)

    with tempfile.TemporaryDirectory() as tmp:
        chart_path = os.path.join(tmp, "chart.png")
        xlsx_path = os.path.join(tmp, f"cleanup_report_{args.date}.xlsx")
        docx_path = os.path.join(tmp, f"cleanup_report_{args.date}.docx")

        build_chart(changelog, chart_path)
        build_excel(changelog, duplicate_pairs, unverified, xlsx_path)
        build_docx(changelog, duplicate_pairs, chart_path, docx_path)
        print(f"Report built: {xlsx_path}, {docx_path}")

        if args.dry_run:
            out_dir = "/tmp/weekly_report_dry_run"
            os.makedirs(out_dir, exist_ok=True)
            shutil.copy(chart_path, os.path.join(out_dir, "chart.png"))
            shutil.copy(xlsx_path, out_dir)
            shutil.copy(docx_path, out_dir)
            print(f"Dry run: saved report files to {out_dir} (no email sent)")
            return

        sender = os.environ["REPORT_SENDER_EMAIL"]
        recipients = [r.strip() for r in os.environ["REPORT_RECIPIENTS"].split(",") if r.strip()]
        send_report_email(xlsx_path, docx_path, changelog, duplicate_pairs, unverified, args.date, sender, recipients)


if __name__ == "__main__":
    main()
